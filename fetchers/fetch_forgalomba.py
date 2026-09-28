#!/usr/bin/env python3
"""
KSH STADAT 24.2.1.22 (sza0070) letöltése és feldolgozása: Magyarországon
első alkalommal forgalomba helyezett személygépkocsik száma gyártmány
szerint, negyedévenként (2022 Q1-től).

Forrás: https://www.ksh.hu/stadat_files/sza/hu/sza0070.html

Ez a "2. lépés" gyors, előretekintő kiegészítője a lassabban mozgó
állomány-adathoz (sza0025) képest: a friss forgalomba helyezés azonnal
mutatja, merre fordul a piac, míg az állomány csak lassan követi.
Fontos: ez a tábla NEM tartalmaz hajtástípus-bontást, csak márkánkénti
bontást (ezért a brand_tier momentumhoz használjuk, nem a drivetrainhez).

Használat:
    python3 fetchers/fetch_forgalomba.py
Kimenet:
    data/forgalomba.json
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import openpyxl

SOURCE_URL = "https://www.ksh.hu/stadat_files/sza/hu/sza0070.xlsx"
SOURCE_PAGE = "https://www.ksh.hu/stadat_files/sza/hu/sza0070.html"
DATA_DIR = Path(__file__).parent.parent / "data"
RAW_PATH = DATA_DIR / "_raw_sza0070.xlsx"
OUT_PATH = DATA_DIR / "forgalomba.json"

TOTAL_LABEL = "Összesen"
SEPARATOR_LABEL = "Ebből:"

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"
)


def download():
    DATA_DIR.mkdir(exist_ok=True)
    subprocess.run(
        [
            "curl", "-sL", "-o", str(RAW_PATH), SOURCE_URL,
            "-A", UA,
            "-H", f"Referer: {SOURCE_PAGE}",
            "-H", "Accept-Language: hu-HU,hu;q=0.9",
        ],
        check=True,
    )


def parse():
    wb = openpyxl.load_workbook(RAW_PATH, data_only=True)
    sheet_name = next(n for n in wb.sheetnames if n != "Információk")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    header_row = rows[1]
    all_quarters = [q for q in header_row[1:] if q is not None]

    # az "Összesen" sor alapján állapítjuk meg, meddig van tényleges adat
    # (a jövőbeli, még üres negyedéveket levágjuk)
    total_row = next(row for row in rows[2:] if row[0] == TOTAL_LABEL)
    total_full = list(total_row[1 : 1 + len(all_quarters)])
    n_valid = len(all_quarters)
    while n_valid > 0 and total_full[n_valid - 1] is None:
        n_valid -= 1

    quarters = all_quarters[:n_valid]
    total = total_full[:n_valid]

    brands: dict[str, list] = {}
    for row in rows[2:]:
        label = row[0]
        if label is None or label == TOTAL_LABEL:
            continue
        if isinstance(label, str) and label.strip().startswith(SEPARATOR_LABEL):
            continue
        values = list(row[1 : 1 + n_valid])
        if any(v is None for v in values):
            continue  # hiányos sor (pl. összevont/törölt márka), kihagyjuk
        brands[label] = values
    return {
        "source": SOURCE_URL,
        "source_page": SOURCE_PAGE,
        "quarters": quarters,
        "total": total,
        "brands": brands,
    }


def main():
    download()
    data = parse()
    OUT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Elmentve: {OUT_PATH} ({len(data['quarters'])} negyedév, {len(data['brands'])} márka)")


if __name__ == "__main__":
    main()

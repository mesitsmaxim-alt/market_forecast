#!/usr/bin/env python3
"""
KSH STADAT 21.1.1.35 (gdp0035) letöltése és feldolgozása: reáljövedelem és
reálkereset index, éves bontásban, "előző év = 100%" bázison.

Forrás: https://www.ksh.hu/stadat_files/gdp/hu/gdp0035.html

A vásárlóerő (reáljövedelem) alakulása a keresleti oldal egyik legfontosabb
makró-mutatója: közvetlenül befolyásolja, hogy a piac az olcsóbb/kényszerű
vagy a drágább/halasztható szegmensek felé tolódik.

Használat:
    python3 fetchers/fetch_realjovedelem.py
Kimenet:
    data/realjovedelem.json
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import openpyxl

SOURCE_URL = "https://www.ksh.hu/stadat_files/gdp/hu/gdp0035.xlsx"
SOURCE_PAGE = "https://www.ksh.hu/stadat_files/gdp/hu/gdp0035.html"
DATA_DIR = Path(__file__).parent.parent / "data"
RAW_PATH = DATA_DIR / "_raw_gdp0035.xlsx"
OUT_PATH = DATA_DIR / "realjovedelem.json"

PREV_YEAR_BASE_LABEL = "Előző év = 100,0%"

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"
)


def download():
    DATA_DIR.mkdir(exist_ok=True)
    subprocess.run(
        [
            "curl", "-sL", "-o", str(RAW_PATH), SOURCE_URL,
            "-A", UA,
            "-H", f"Referer: {SOURCE_PAGE}",
            "-H", "Accept-Language: hu-HU,hu;q=0.9",
        ],
        check=True,
    )


def parse():
    wb = openpyxl.load_workbook(RAW_PATH, data_only=True)
    sheet_name = next(n for n in wb.sheetnames if n != "Információk")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    # a fájl két szekciót tartalmaz (1990=100% kumulált, majd előző év=100%
    # YoY bázis) - nekünk az utóbbi kell, ami a "Előző év = 100,0%" felirat
    # utáni sorokban van
    start = next(i for i, r in enumerate(rows) if r[0] == PREV_YEAR_BASE_LABEL) + 1

    years, wage_yoy_pct, income_yoy_pct = [], [], []
    for row in rows[start:]:
        year, wage_idx, _dummy, income_idx = row[0], row[2], None, row[3]
        if year is None:
            continue
        years.append(int(year))
        wage_yoy_pct.append(round(wage_idx - 100, 2) if isinstance(wage_idx, (int, float)) else None)
        income_yoy_pct.append(round(income_idx - 100, 2) if isinstance(income_idx, (int, float)) else None)

    return {
        "source": SOURCE_URL,
        "source_page": SOURCE_PAGE,
        "note": (
            "reálkereset = egy keresőre jutó reálbér, reáljövedelem = egy "
            "főre jutó teljes reáljövedelem (szélesebb, a vásárlóerő jobb "
            "proxyja). Az értékek YoY %-os változást jelentenek."
        ),
        "years": years,
        "real_wage_yoy_pct": wage_yoy_pct,
        "real_income_yoy_pct": income_yoy_pct,
    }


def main():
    download()
    data = parse()
    OUT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Elmentve: {OUT_PATH} ({data['years'][0]}–{data['years'][-1]})")


if __name__ == "__main__":
    main()
